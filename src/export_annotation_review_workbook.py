"""Export a compact, single-sheet annotation review workbook."""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HEADERS = [
    "Review order",
    "Priority",
    "Review focus",
    "Raw CellTypist label",
    "Proposed broad population",
    "Proposed use",
    "Cells",
    "Subjects",
    "Samples",
    "Batches",
    "Median confidence",
    "Marker panel",
    "Expected-marker detection",
    "Cluster summary",
    "Batch summary",
    "Method eligibility",
    "Marker / PBMC assessment",
    "Cluster assessment",
    "Study / technical support",
    "Final decision",
    "Final analysis label",
    "Final use",
    "Evidence-based rationale",
    "Review status",
]

PRIORITY = {
    "primary": "1 - Primary",
    "exploratory": "2 - Exploratory",
    "unresolved": "3 - Unresolved",
}

FOCUS = {
    "primary": "Confirm broad mapping for primary inference",
    "exploratory": "Confirm exclusion from primary inference",
    "unresolved": "Confirm that no defensible broad parent is supported",
}


def _number(value: str, default: float = 0.0) -> float:
    if value in {"", "NA", "NaN", "nan"}:
        return default
    return float(value)


def _integer(value: str) -> int:
    return int(_number(value))


def _yes_no(value: str) -> str:
    return "yes" if str(value).strip().lower() == "true" else "no"


def _sort_key(row: dict[str, str]) -> tuple[int, str, int, str]:
    disposition = row["disposition"]
    rank = {"primary": 1, "exploratory": 2, "unresolved": 3}[disposition]
    broad_label = row["analysis_label"] or "ZZZ unresolved"
    if disposition == "primary":
        return rank, broad_label, -_integer(row["n_cells"]), row["raw_label"]
    return rank, "", -_integer(row["n_cells"]), row["raw_label"]


def _cluster_summary(row: dict[str, str]) -> str:
    return (
        f"Leiden {row['dominant_leiden']}; "
        f"{_number(row['dominant_leiden_fraction']):.1%} of label; "
        f"{_integer(row['n_leiden_clusters'])} clusters"
    )


def _batch_summary(row: dict[str, str]) -> str:
    return (
        f"{row['dominant_batch']}; "
        f"{_number(row['dominant_batch_fraction']):.1%} of label; "
        f"{_integer(row['n_batches_detected'])} batches"
    )


def _method_summary(row: dict[str, str]) -> str:
    methods = (
        ("Composition", "composition_passes_current_contract"),
        ("Signature", "signature_passes_current_contract"),
        ("Pseudobulk", "pseudobulk_passes_current_contract"),
        ("Prediction", "age_prediction_figure_passes_current_contract"),
    )
    return " | ".join(f"{name}: {_yes_no(row[column])}" for name, column in methods)


def _read_rows(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    required = {
        "raw_label",
        "analysis_label",
        "disposition",
        "n_cells",
        "n_subjects_detected",
        "n_sample_units_detected",
        "n_batches_detected",
        "confidence_median",
        "marker_panel_label",
        "expected_marker_detection",
        "dominant_leiden",
        "dominant_leiden_fraction",
        "n_leiden_clusters",
        "dominant_batch",
        "dominant_batch_fraction",
        "composition_passes_current_contract",
        "signature_passes_current_contract",
        "pseudobulk_passes_current_contract",
        "age_prediction_figure_passes_current_contract",
    }
    missing = sorted(required - set(rows[0] if rows else []))
    if missing:
        raise ValueError(f"Review table is missing required columns: {missing}")
    if len(rows) != 84 or len({row["raw_label"] for row in rows}) != 84:
        raise ValueError("Expected exactly 84 unique raw labels")
    return sorted(rows, key=_sort_key)


def _add_validation(sheet, cell_range: str, values: list[str]) -> None:
    validation = DataValidation(
        type="list",
        formula1='"' + ",".join(values) + '"',
        allow_blank=True,
    )
    validation.error = "Select one of the predefined review values."
    validation.errorTitle = "Invalid review value"
    validation.prompt = "Choose a value from the list."
    validation.promptTitle = "Human-review field"
    validation.showErrorMessage = True
    validation.showInputMessage = True
    sheet.add_data_validation(validation)
    validation.add(cell_range)


def export_workbook(review_table: Path, output_path: Path) -> None:
    rows = _read_rows(review_table)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Annotation review"
    sheet.sheet_view.showGridLines = False

    sheet.merge_cells("A1:X1")
    sheet["A1"] = "Blood Age Atlas annotation review"
    sheet["A1"].font = Font(name="Aptos Display", size=16, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="174A5B")
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 30

    sheet.merge_cells("A2:X2")
    sheet["A2"] = (
        "Purpose: decide analytical treatment of each raw CellTypist label; "
        "this is not proof of exact biological identity."
    )
    sheet["A2"].font = Font(name="Aptos", size=10, italic=True, color="2F3E46")
    sheet["A2"].fill = PatternFill("solid", fgColor="E8F1F2")
    sheet["A2"].alignment = Alignment(vertical="center")

    sheet.merge_cells("A3:X3")
    sheet["A3"] = (
        "Start with priority 1 (17 primary proposals). Evidence columns are fixed; "
        "complete only the yellow review columns. Do not inspect age-association outcomes."
    )
    sheet["A3"].font = Font(name="Aptos", size=10, bold=True, color="6B4F00")
    sheet["A3"].fill = PatternFill("solid", fgColor="FFF2CC")
    sheet["A3"].alignment = Alignment(vertical="center")

    sheet.merge_cells("A4:X4")
    sheet["A4"] = (
        "Decision rule: accept a broad mapping only when marker/PBMC plausibility, cluster coherence, "
        "and study/technical support show no important contradiction. No universal threshold is implied."
    )
    sheet["A4"].font = Font(name="Aptos", size=9, color="37474F")
    sheet["A4"].fill = PatternFill("solid", fgColor="F4F6F6")
    sheet["A4"].alignment = Alignment(vertical="center")

    header_row = 6
    for column, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(header_row, column, header)
        cell.font = Font(name="Aptos", size=9, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="276678")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[header_row].height = 42

    editable_fill = PatternFill("solid", fgColor="FFF2CC")
    evidence_fill = PatternFill("solid", fgColor="F7F9FA")
    primary_fill = PatternFill("solid", fgColor="DDEBF7")
    exploratory_fill = PatternFill("solid", fgColor="E2F0D9")
    unresolved_fill = PatternFill("solid", fgColor="FCE4D6")
    thin_bottom = Border(bottom=Side(style="thin", color="D9E1E5"))

    first_data_row = header_row + 1
    for index, row in enumerate(rows, start=first_data_row):
        disposition = row["disposition"]
        values = [
            index - header_row,
            PRIORITY[disposition],
            FOCUS[disposition],
            row["raw_label"],
            row["analysis_label"],
            disposition,
            _integer(row["n_cells"]),
            _integer(row["n_subjects_detected"]),
            _integer(row["n_sample_units_detected"]),
            _integer(row["n_batches_detected"]),
            _number(row["confidence_median"]),
            row["marker_panel_label"],
            row["expected_marker_detection"],
            _cluster_summary(row),
            _batch_summary(row),
            _method_summary(row),
            "",
            "",
            "",
            "",
            row["analysis_label"],
            disposition,
            "",
            f'=IF(AND(Q{index}<>"",R{index}<>"",S{index}<>"",T{index}<>"",V{index}<>"",W{index}<>""),"Complete","Pending")',
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(index, column, value)
            cell.font = Font(name="Aptos", size=9, color="1F2933")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = thin_bottom
            if 17 <= column <= 23:
                cell.fill = editable_fill
            else:
                cell.fill = evidence_fill
        sheet.cell(index, 2).fill = {
            "primary": primary_fill,
            "exploratory": exploratory_fill,
            "unresolved": unresolved_fill,
        }[disposition]
        sheet.cell(index, 2).font = Font(name="Aptos", size=9, bold=True, color="1F2933")
        sheet.row_dimensions[index].height = 34

    last_row = first_data_row + len(rows) - 1
    sheet.auto_filter.ref = f"A{header_row}:X{last_row}"
    sheet.freeze_panes = "G7"

    _add_validation(
        sheet,
        f"Q{first_data_row}:Q{last_row}",
        ["coherent", "inconclusive", "incoherent", "not_assessable"],
    )
    _add_validation(
        sheet,
        f"R{first_data_row}:R{last_row}",
        ["coherent", "inconclusive", "incoherent", "not_assessable"],
    )
    _add_validation(
        sheet,
        f"S{first_data_row}:S{last_row}",
        ["adequate", "inconclusive", "inadequate", "not_assessable"],
    )
    _add_validation(
        sheet,
        f"T{first_data_row}:T{last_row}",
        ["accept_as_proposed", "modify", "reject_pending_resolution"],
    )
    _add_validation(
        sheet,
        f"V{first_data_row}:V{last_row}",
        ["primary", "exploratory", "unresolved"],
    )

    pending_fill = PatternFill("solid", fgColor="F4CCCC")
    complete_fill = PatternFill("solid", fgColor="D9EAD3")
    sheet.conditional_formatting.add(
        f"X{first_data_row}:X{last_row}",
        FormulaRule(formula=[f'X{first_data_row}="Pending"'], fill=pending_fill),
    )
    sheet.conditional_formatting.add(
        f"X{first_data_row}:X{last_row}",
        FormulaRule(formula=[f'X{first_data_row}="Complete"'], fill=complete_fill),
    )

    widths = {
        "A": 10,
        "B": 16,
        "C": 31,
        "D": 32,
        "E": 25,
        "F": 13,
        "G": 11,
        "H": 10,
        "I": 10,
        "J": 9,
        "K": 14,
        "L": 24,
        "M": 48,
        "N": 31,
        "O": 30,
        "P": 48,
        "Q": 24,
        "R": 21,
        "S": 25,
        "T": 27,
        "U": 25,
        "V": 15,
        "W": 52,
        "X": 14,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    for column in ("A", "B", "F", "G", "H", "I", "J", "K", "X"):
        for cell in sheet[column][header_row:last_row]:
            cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
    for cell in sheet["K"][first_data_row - 1 : last_row]:
        cell.number_format = "0.000"
    for column in ("G", "H", "I", "J"):
        for cell in sheet[column][first_data_row - 1 : last_row]:
            cell.number_format = "#,##0"

    sheet.print_title_rows = f"1:{header_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.auto_filter.ref = f"A{header_row}:X{last_row}"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)

    verification = load_workbook(output_path, data_only=False)
    verified_sheet = verification["Annotation review"]
    if verified_sheet.max_row != 90 or verified_sheet.max_column != 24:
        raise RuntimeError("Workbook dimensions do not match the 84-label review contract")
    priorities = [verified_sheet.cell(row, 2).value for row in range(7, 91)]
    if priorities[:17] != ["1 - Primary"] * 17:
        raise RuntimeError("The first 17 workbook rows are not all primary proposals")
    if verified_sheet.auto_filter.ref != "A6:X90":
        raise RuntimeError("Workbook filter range is missing or incorrect")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--review-table",
        type=Path,
        default=Path(
            "results/blood_age_atlas_requalification_longitudinal/validation/"
            "annotation_mapping_review_table.csv"
        ),
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("docs/reviews/annotation_mapping_review.xlsx"),
    )
    return parser.parse_args()


if __name__ == "__main__":
    arguments = parse_args()
    export_workbook(arguments.review_table, arguments.output)
